"""Exportação do relatório (Dossier de recette) em Excel, com a mesma estrutura do FiberCable."""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style='thin', color='808080')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
SUB_FILL = PatternFill('solid', fgColor='D9E2F3')
OUT_FILL = PatternFill('solid', fgColor='FFC7CE')
OUT_FONT = Font(bold=True, color='C00000', size=10)
CALC_FONT = Font(italic=True, color='1F4E79', size=10)
MISS_FILL = PatternFill('solid', fgColor='FFF2CC')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
NUM3 = '0.000'
NUM2 = '0.00'


def _cell(ws, r, c, v=None, bold=False, fill=None, fmt=None, border=True, color=None, size=None, align=CENTER):
    cell = ws.cell(row=r, column=c, value=v)
    if bold or color or size:
        cell.font = Font(bold=bold, color=color, size=size or 10)
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = BORDER
    if align:
        cell.alignment = align
    return cell


def _merge(ws, r1, c1, r2, c2, v=None, **kw):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            _cell(ws, r, c, **{k: x for k, x in kw.items() if k in ('fill', 'border')})
    _cell(ws, r1, c1, v, **kw)
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _header(ws, meta, title, width):
    _merge(ws, 1, 2, 1, width, meta['address'], bold=True, border=False, size=12)
    _cell(ws, 3, width, meta['cable'], border=False, align=Alignment(horizontal='right'))
    _cell(ws, 4, width, meta['origin'], border=False, align=Alignment(horizontal='right'))
    _cell(ws, 5, width, meta['end'], border=False, align=Alignment(horizontal='right'))
    _merge(ws, 7, 2, 7, width, title, bold=True, fill=HEAD_FILL, color='FFFFFF', size=12)


def _val(ws, r, c, v, fmt=NUM3, out=False, missing=False, calc=False):
    """Valor numérico: a vermelho quando ultrapassa o limiar, itálico azul quando calculado na curva."""
    fill = OUT_FILL if out else (MISS_FILL if missing else None)
    cell = _cell(ws, r, c, v, fmt=fmt, fill=fill)
    if out:
        cell.font = OUT_FONT
    elif calc:
        cell.font = CALC_FONT
    return cell


def build_workbook(res, meta):
    meta = dict(meta)
    info = res['info']
    meta.setdefault('cable', info['cable'])
    meta.setdefault('origin', info['origin'])
    meta.setdefault('end', info['end'])
    meta.setdefault('address', meta.get('end') or '')
    fibers, wls = res['fibers'], res['wls']
    points, sections = res['points'], res['sections']
    th = res['thresholds']

    wb = Workbook()

    # --- Titre
    ws = wb.active
    ws.title = 'Titre'
    ws.column_dimensions['B'].width = 4
    _merge(ws, 4, 2, 4, 7, 'Dossier de recette du câble :', bold=True, size=16, border=False)
    _cell(ws, 6, 4, 'Nom du câble :', bold=True, border=False)
    _cell(ws, 6, 5, meta['cable'], bold=True, size=14, border=False)
    _cell(ws, 12, 3, 'Chef de projet', border=False); _cell(ws, 12, 5, meta.get('manager', ''), border=False)
    _cell(ws, 13, 3, 'Raccordement', border=False); _cell(ws, 13, 5, meta.get('company', ''), border=False)
    _cell(ws, 14, 3, 'Le :', border=False)
    d = _cell(ws, 14, 6, meta.get('date') or datetime.date.today(), border=False)
    d.number_format = 'DD/MM/YYYY'
    _cell(ws, 15, 3, 'Signature :', border=False)
    for c in 'CDEF':
        ws.column_dimensions[c].width = 18

    # --- Table des matières
    ws = wb.create_sheet('Table Matières')
    _header(ws, meta, 'Contenu du rapport', 5)
    for i, (t, n) in enumerate([('Informations & Généralités', 1), ('Informations mesures', 2),
                                ('Connecteurs & épissures', 3), ('Tableau des réflectances', 4),
                                ('Pentes', 5), ('Valeurs hors normes', 6)]):
        _cell(ws, 10 + i, 2, t, border=False, align=Alignment(horizontal='left'))
        _cell(ws, 10 + i, 4, n, border=False)
    ws.column_dimensions['B'].width = 32
    ws.column_dimensions['E'].width = 30

    # --- Informations générales
    ws = wb.create_sheet('Informations Générales')
    _header(ws, meta, 'Informations & Généralités', 7)
    left = Alignment(horizontal='left')
    rows = [(9, 2, 'Informations concernant la liaison', True),
            (11, 3, 'Nom du câble :', False), (12, 3, 'Origine :', False), (13, 3, 'Extrémité :', False),
            (16, 2, 'Informations concernant le câble', True),
            (19, 3, 'Raccordement', False), (20, 3, 'Fabricant du câble', False), (21, 3, 'Type de câble', False),
            (22, 3, 'Fabricant de la fibre', False), (23, 3, 'Type de fibre', False),
            (25, 3, 'Nombre de fibres :', False), (26, 3, 'Numéro première fibre :', False),
            (27, 3, 'Numéro dernière fibre :', False), (28, 3, 'Nombre de longueurs :', False)]
    for r, c, t, b in rows:
        _cell(ws, r, c, t, bold=b, border=False, align=left)
    vals = {11: meta['cable'], 12: meta['origin'], 13: meta['end'], 19: meta.get('company', ''),
            20: meta.get('cable_maker', ''), 21: meta.get('cable_type', ''), 22: meta.get('fiber_maker', ''),
            23: meta.get('fiber_type', ''), 25: len(fibers), 26: min(fibers), 27: max(fibers),
            28: len(sections)}
    for r, v in vals.items():
        _cell(ws, r, 5, v, border=False, align=left)
    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['E'].width = 32

    # --- Informations mesure
    ws = wb.create_sheet('Informations Mesure')
    _header(ws, meta, 'Informations mesures', 10)
    _cell(ws, 11, 2, 'OTDR', bold=True, border=False, align=left)
    _cell(ws, 12, 6, 'Origine', bold=True); _cell(ws, 12, 8, 'Extrémité', bold=True)
    lines = [('Type de réflectomètre :', info['otdr']), ('N° série du châssis :', info['otdr_sn']),
             ('Type de module :', info['module']), ('N° série du module :', info['module_sn'])]
    for i, (t, v) in enumerate(lines):
        _cell(ws, 13 + i, 3, t, border=False, align=left)
        _cell(ws, 13 + i, 6, v); _cell(ws, 13 + i, 8, v)
    _cell(ws, 18, 3, "Longueurs d'onde réelles (nm) :", border=False, align=left)
    pw = info['per_wl']
    for j, wl in enumerate(wls):
        c = 8 + j
        _cell(ws, 18, c, pw[wl]['wavelength'], bold=True)
        _cell(ws, 19, c, pw[wl]['ior'], fmt='0.0000')
        _cell(ws, 20, c, pw[wl]['pulse_ns'])
        _cell(ws, 22, c, th['splice'], fmt=NUM2)
        _cell(ws, 23, c, th['splice'], fmt=NUM2)
        _cell(ws, 24, c, th['connector'], fmt=NUM2)
        _cell(ws, 25, c, th['connector'], fmt=NUM2)
        _cell(ws, 26, c, th['slope'].get(wl), fmt=NUM2)
        _cell(ws, 27, c, th['slope'].get(wl), fmt=NUM2)
        _cell(ws, 28, c, th['reflectance'])
        for r in (29, 30, 31):
            _cell(ws, r, c, '- - -')
        _cell(ws, 34, c, pw[wl]['backscatter'])
    for r, t in [(19, 'Indice de la fibre :'), (20, 'Impulsion (ns) :'), (21, 'Valeurs des seuils (dB) :'),
                 (22, '   Epissure :'), (23, '   Moyenne épissures :'), (24, '   Connecteur :'),
                 (25, '   Moyenne connecteurs :'), (26, '   Pente :'), (27, '   Moyenne pentes :'),
                 (28, '   Réflectance :'), (29, '   Bilan min :'), (30, '   Bilan max :'), (31, '   ORL max :'),
                 (32, '   Seuil écart 1310/1550nm (dB) :'), (34, 'Coefficient de rétrodiffusion (dB) :')]:
        _cell(ws, r, 3, t, border=False, align=left)
    _cell(ws, 32, 5, th['diff_1550_1310'], fmt=NUM2)
    _cell(ws, 39, 2, 'Logiciels', bold=True, border=False, align=left)
    _cell(ws, 41, 3, 'APP FIBER CABLE (protótipo)', border=False, align=left)
    _cell(ws, 42, 3, 'Logiciel OTDR', border=False, align=left)
    _cell(ws, 42, 6, info['software'], border=False, align=left)
    ws.column_dimensions['C'].width = 34

    # --- Epissures
    ws = wb.create_sheet('Epissures')
    has_diff = 1310 in wls and 1550 in wls
    block = 3 * len(wls) + (1 if has_diff else 0)
    ncols = 1 + block * len(points) + 3 * len(wls)
    _header(ws, meta, 'Connecteurs & épissures', ncols)
    r0 = 9
    _cell(ws, r0, 1, 'Evénement', bold=True, fill=SUB_FILL)
    _cell(ws, r0 + 1, 1, 'Long. Cum. (m)', bold=True, fill=SUB_FILL)
    _cell(ws, r0 + 2, 1, 'Long (m)', bold=True, fill=SUB_FILL)
    _merge(ws, r0 + 3, 1, r0 + 4, 1, 'Numéro fibre', bold=True, fill=SUB_FILL)
    for i, p in enumerate(points):
        c0 = 2 + i * block
        _merge(ws, r0, c0, r0, c0 + block - 1, p['name'], bold=True, fill=SUB_FILL)
        _merge(ws, r0 + 1, c0, r0 + 1, c0 + block - 1, round(p['cum_m'], 2), fmt=NUM2)
        _merge(ws, r0 + 2, c0, r0 + 2, c0 + block - 1, round(p['len_m'], 2), fmt=NUM2)
        for j, wl in enumerate(wls):
            cc = c0 + 3 * j
            _merge(ws, r0 + 3, cc, r0 + 3, cc + 2, f'{wl}nm', bold=True, fill=SUB_FILL)
            for k, t in enumerate(('OE', 'EO', 'Moyenne')):
                _cell(ws, r0 + 4, cc + k, t, bold=True, fill=SUB_FILL)
        if has_diff:
            _merge(ws, r0 + 3, c0 + block - 1, r0 + 4, c0 + block - 1, '1550-1310 Différence', bold=True, fill=SUB_FILL)
    cs = 2 + block * len(points)
    for k, t in enumerate(('Moyenne', 'Maximum', 'Bilan de liaison')):
        cc = cs + k * len(wls)
        _merge(ws, r0, cc, r0 + 3, cc + len(wls) - 1, t, bold=True, fill=SUB_FILL)
        for j, wl in enumerate(wls):
            _cell(ws, r0 + 4, cc + j, f'{wl}nm', bold=True, fill=SUB_FILL)
    r = r0 + 5
    for f in fibers:
        _cell(ws, r, 1, f, bold=True)
        for i, p in enumerate(points):
            c0 = 2 + i * block
            for j, wl in enumerate(wls):
                d = p['data'][f'{f}|{wl}']
                cc = c0 + 3 * j
                _val(ws, r, cc, d['OE'], missing=d['OE_missing'], out=d.get('OE_out'), calc=d.get('OE_calc'))
                _val(ws, r, cc + 1, d['EO'], missing=d['EO_missing'], out=d.get('EO_out'), calc=d.get('EO_calc'))
                _val(ws, r, cc + 2, d['avg'], out=d['out'])
            if has_diff:
                dv = p['data'].get(f'{f}|diff')
                _val(ws, r, c0 + block - 1, dv, out=dv is not None and dv > th['diff_1550_1310'])
        for j, wl in enumerate(wls):
            s = res['summary'][f'{f}|{wl}']
            _val(ws, r, cs + j, s['splice_mean'], fmt='0.0000')
            _val(ws, r, cs + len(wls) + j, s['splice_max'])
            _val(ws, r, cs + 2 * len(wls) + j, s['budget'])
        r += 1
    # linha de médias por coluna
    for i, p in enumerate(points):
        c0 = 2 + i * block
        for j, wl in enumerate(wls):
            vals = [p['data'][f'{f}|{wl}']['avg'] for f in fibers if p['data'][f'{f}|{wl}']['avg'] is not None]
            if vals:
                _cell(ws, r, c0 + 3 * j + 2, sum(vals) / len(vals), fmt='0.0000', bold=True)
    if any(p['data'][f'{f}|{wl}'].get(d + '_calc') for p in points for f in fibers for wl in wls for d in ('OE', 'EO')):
        note = _cell(ws, r + 2, 1, 'Valeurs en italique bleu : événement non détecté par le réflectomètre, '
                     'perte calculée sur la courbe (méthode 2 points LSA).', border=False, align=Alignment(horizontal='left'))
        note.font = CALC_FONT
    ws.column_dimensions['A'].width = 15
    for c in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9

    # --- Réflectances
    ws = wb.create_sheet('Réflectances')
    conn = [p for p in points if p['kind'] == 'connector']
    ncol = 1 + 2 * len(wls) * len(conn) + 1
    _header(ws, meta, 'Tableau des réflectances', ncol)
    _cell(ws, r0, 1, 'Evénement', bold=True, fill=SUB_FILL)
    _merge(ws, r0 + 1, 1, r0 + 2, 1, 'Numéro fibre', bold=True, fill=SUB_FILL)
    labels = ['Connecteur O', 'Connecteur E']
    for i, p in enumerate(conn):
        c0 = 2 + i * 2 * len(wls)
        _merge(ws, r0, c0, r0, c0 + 2 * len(wls) - 1, labels[i] if i < 2 else p['name'], bold=True, fill=SUB_FILL)
        for j, wl in enumerate(wls):
            _merge(ws, r0 + 1, c0 + 2 * j, r0 + 1, c0 + 2 * j + 1, f'{wl}nm', bold=True, fill=SUB_FILL)
            _cell(ws, r0 + 2, c0 + 2 * j, 'OE', bold=True, fill=SUB_FILL)
            _cell(ws, r0 + 2, c0 + 2 * j + 1, 'EO', bold=True, fill=SUB_FILL)
    _merge(ws, r0, ncol, r0 + 2, ncol, 'Commentaire', bold=True, fill=SUB_FILL)
    r = r0 + 3
    for f in fibers:
        _cell(ws, r, 1, f, bold=True)
        for i, p in enumerate(conn):
            c0 = 2 + i * 2 * len(wls)
            for j, wl in enumerate(wls):
                d = p['data'][f'{f}|{wl}']
                _val(ws, r, c0 + 2 * j, d['OE_refl'], fmt=NUM2, out=d.get('OE_refl_out'))
                _val(ws, r, c0 + 2 * j + 1, d['EO_refl'], fmt=NUM2, out=d.get('EO_refl_out'))
        _cell(ws, r, ncol, '')
        r += 1
    for c in range(2, ncol):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions[get_column_letter(ncol)].width = 30

    # --- Pentes
    ws = wb.create_sheet('Pentes')
    block = 3 * len(wls)
    ncol = 1 + block * len(sections) + len(wls)
    _header(ws, meta, 'Pentes', ncol)
    _cell(ws, r0, 1, 'Evénement', bold=True, fill=SUB_FILL)
    _cell(ws, r0 + 1, 1, 'Long (m)', bold=True, fill=SUB_FILL)
    _merge(ws, r0 + 2, 1, r0 + 3, 1, 'Numéro fibre', bold=True, fill=SUB_FILL)
    for i, s in enumerate(sections):
        c0 = 2 + i * block
        _merge(ws, r0, c0, r0, c0 + block - 1, f"Pente N° {s['index']}", bold=True, fill=SUB_FILL)
        _merge(ws, r0 + 1, c0, r0 + 1, c0 + block - 1, round(s['len_m'], 2), fmt=NUM2)
        for j, wl in enumerate(wls):
            _merge(ws, r0 + 2, c0 + 3 * j, r0 + 2, c0 + 3 * j + 2, f'{wl}nm', bold=True, fill=SUB_FILL)
            for k, t in enumerate(('OE', 'EO', 'Moyenne')):
                _cell(ws, r0 + 3, c0 + 3 * j + k, t, bold=True, fill=SUB_FILL)
    cs = 2 + block * len(sections)
    _merge(ws, r0, cs, r0 + 2, cs + len(wls) - 1, 'Moyenne pentes', bold=True, fill=SUB_FILL)
    for j, wl in enumerate(wls):
        _cell(ws, r0 + 3, cs + j, f'{wl}nm', bold=True, fill=SUB_FILL)
    r = r0 + 4
    for f in fibers:
        _cell(ws, r, 1, f, bold=True)
        for i, s in enumerate(sections):
            c0 = 2 + i * block
            for j, wl in enumerate(wls):
                d = s['data'][f'{f}|{wl}']
                _val(ws, r, c0 + 3 * j, d['OE'], out=d.get('OE_out'))
                _val(ws, r, c0 + 3 * j + 1, d['EO'], out=d.get('EO_out'))
                _val(ws, r, c0 + 3 * j + 2, d['avg'], out=d['out'])
        for j, wl in enumerate(wls):
            _val(ws, r, cs + j, res['summary'][f'{f}|{wl}']['slope_mean'])
        r += 1
    ws.column_dimensions['A'].width = 15
    for c in range(2, ncol + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9

    # --- Hors normes
    ws = wb.create_sheet('Hors Normes')
    _header(ws, meta, 'Valeurs hors normes', 10)
    out = res['out']
    r = 9

    def table(title, heads, rows):
        nonlocal r
        _merge(ws, r, 2, r, 10, title, bold=True, fill=SUB_FILL)
        r += 2
        for k, h in enumerate(heads):
            _cell(ws, r, 2 + k, h, bold=True, fill=SUB_FILL)
        _merge(ws, r, 2 + len(heads), r, 10, 'Commentaire', bold=True, fill=SUB_FILL)
        r += 1
        if not rows:
            _merge(ws, r, 2, r, 10, 'Pas de valeurs')
            r += 1
        for row in rows:
            for k, v in enumerate(row):
                cell = _cell(ws, r, 2 + k, v, fmt=NUM3 if isinstance(v, float) else None)
                if k == 2:  # o valor fora do limiar
                    cell.font = OUT_FONT
            _merge(ws, r, 2 + len(row), r, 10, '')
            r += 1
        r += 1

    table('Tableau des connecteurs et épissures',
          ['N° Epissure', 'N° Fibre', 'Moyenne (dB)', "Longueur d'onde (nm)"],
          [(o['point'], o['fiber'], o['value'], o['wl']) for o in out['splices']])
    table('Tableau des pentes', ['N° Pente', 'N° Fibre', 'Moyenne (dB/km)', "Longueur d'onde (nm)"],
          [(o['section'], o['fiber'], o['value'], o['wl']) for o in out['slopes']])
    table('Tableau des réflectances', ['Réflectance', 'N° Fibre', 'Valeur (dB)', "Longueur d'onde (nm)"],
          [(o['point'], o['fiber'], o['value'], o['wl']) for o in out['refl']])
    table('Tableau des écarts 1550-1310', ['Evénement', 'N° Fibre', 'Ecart (dB)'],
          [(o['point'], o['fiber'], o['value']) for o in out['diff']])
    table('Tableau des bilans', ['N° Fibre', 'Moyenne (dB)', "Longueur d'onde (nm)"], [])
    ws.column_dimensions['B'].width = 26
    for c in 'CDE':
        ws.column_dimensions[c].width = 16

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = False
        sh.page_setup.orientation = 'landscape'
        sh.page_setup.fitToWidth = 1
        sh.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
